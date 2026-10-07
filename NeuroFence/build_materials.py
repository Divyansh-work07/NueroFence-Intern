from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC

ROOT = Path(__file__).resolve().parent
for folder in ('docs', 'presentation'):
    (ROOT/folder).mkdir(parents=True, exist_ok=True)
source = (ROOT/'docs/project_report.md').read_text(encoding='utf-8')
doc = Document(); sec = doc.sections[0]; sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7)
doc.styles['Normal'].font.name = 'Aptos'; doc.styles['Normal'].font.size = Pt(10)
for line in source.splitlines():
    s = line.strip()
    if not s: continue
    if s.startswith('# '): doc.add_heading(s[2:], 0)
    elif s.startswith('## '): doc.add_heading(s[3:], 1)
    elif s.startswith('- '): doc.add_paragraph(s[2:], style='List Bullet')
    elif len(s)>2 and s[0].isdigit() and s[1:3]=='. ': doc.add_paragraph(s[3:], style='List Number')
    else: doc.add_paragraph(s.replace('**',''))
doc.save(ROOT/'docs/NeuroFence_Project_Report.docx')

prs = Presentation(); prs.slide_width = PI(13.333); prs.slide_height = PI(7.5)
navy=PC(8,14,25); teal=PC(85,224,193); white=PC(230,238,248); muted=PC(173,190,209)
slides = [
('NeuroFence','LLM Weight Poisoning & Backdoor Scanner\nCyber Security Internship Project\n05 Sep – 05 Oct 2026','MODEL ARTIFACT INTAKE SECURITY'),
('The intake problem','Model files can be altered, malformed or difficult to verify.\nUntrusted artifacts should not be executed during initial review.\nA clean static scan cannot prove a backdoor is absent.','WHY THIS PROJECT'),
('What the prototype does','Calculates SHA-256 for every upload.\nInspects safe NPZ and safetensors content.\nFlags non-finite values, unusual magnitudes, archive hazards and selected config keys.\nStores history and produces downloadable JSON reports.','SCOPE'),
('Architecture','Browser dashboard → FastAPI upload API → static scanner → heuristic risk engine\n\nSQLite scan history · JSON report · optional PDF helper\n\nTemporary upload is removed after scanning.','SYSTEM'),
('Static checks','File identity: SHA-256, size and format\nTensor evidence: shapes, dtypes, extrema, mean, standard deviation and zero fraction\nContainer evidence: ZIP member paths and compression ratios\nPickle formats are flagged and never loaded.','ANALYSIS'),
('Risk score','Severity points: LOW 5 · MEDIUM 18 · HIGH 35 · CRITICAL 50\n\nBands: LOW <20 · MEDIUM 20–49 · HIGH 50–74 · CRITICAL ≥75\n\nThe score is a heuristic, not a probability or verdict.','DECISION MODEL'),
('Synthetic demonstration','Clean fixture: finite synthetic arrays; expected LOW risk.\nAnomalous fixture: planted NaN and magnitude 250; expected HIGH risk.\nPickle placeholder: flagged by extension without loading.\n\nThese examples show specific checks, not general backdoor detection.','EVALUATION'),
('What the evidence means','An anomaly is an indicator. Compare it with an authenticated, architecture-matched checkpoint.\n\nGlobal statistics can miss crafted backdoors. Legitimate models can also exceed conservative thresholds.','INTERPRETATION'),
('Safety boundaries','No arbitrary code execution. No pickle deserialization.\nNPZ loaded with allow_pickle=False. Upload limit: 512 MiB.\nUse localhost; isolate parser workloads before broader deployment.','SECURE BY DESIGN'),
('Next steps','Add signed baseline comparison.\nAdd a resource-limited inference sandbox and reviewed behavioral test suite.\nCalibrate thresholds on labeled models.\nAdd authentication, retention controls and deployment hardening.','ROADMAP'),
('Demonstration','1. Start with run_windows.bat.\n2. Open http://127.0.0.1:8000.\n3. Scan demo/clean_demo.npz.\n4. Scan demo/anomalous_demo.npz.\n5. Compare findings and download JSON.','LIVE WALKTHROUGH'),
('Conclusion','NeuroFence provides a safe first-pass model intake workflow with evidence, history and reports.\n\nIt supports triage; it does not certify model safety.','THANK YOU')]
for title, body, kicker in slides:
    sl=prs.slides.add_slide(prs.slide_layouts[6]); sl.background.fill.solid(); sl.background.fill.fore_color.rgb=navy
    box=sl.shapes.add_textbox(PI(.8),PI(.55),PI(11.7),PI(.4)); p=box.text_frame.paragraphs[0]; p.text=kicker; p.font.size=PP(12); p.font.bold=True; p.font.color.rgb=teal
    box=sl.shapes.add_textbox(PI(.8),PI(1.15),PI(11.8),PI(.85)); p=box.text_frame.paragraphs[0]; p.text=title; p.font.size=PP(32); p.font.bold=True; p.font.color.rgb=white
    box=sl.shapes.add_textbox(PI(.9),PI(2.3),PI(11.6),PI(4.6)); tf=box.text_frame; tf.word_wrap=True
    for i, part in enumerate(body.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=part; p.font.size=PP(20 if len(part)<105 else 17); p.font.color.rgb=muted; p.space_after=PP(15)
prs.save(ROOT/'presentation/NeuroFence_Internship_Presentation.pptx')
print('Created DOCX report and PPTX presentation.')
